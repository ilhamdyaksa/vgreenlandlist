import streamlit as st
import pandas as pd
import docx
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
import io
import zipfile
import re
import os
import requests
import urllib.parse
from datetime import datetime

st.set_page_config(page_title="Automasi LOI - MSS Alpha Tech", page_icon="⚡", layout="wide")

TEMPLATE_FILE = "LOI - Contoh Template Sewa Lahan.docx"

def format_rupiah(val):
    try:
        val = float(val)
        if pd.isna(val) or val <= 0:
            return "Rp -"
        return f"Rp {int(val):,.0f}".replace(",", ".")
    except:
        return "Rp -"

def clean_text(val):
    if pd.isna(val) or val is None or str(val).lower() == 'nan':
        return ""
    return str(val).strip()

def extract_price(val, fallback):
    if pd.isna(val) or val is None:
        return float(fallback)
    s = str(val)
    digits = re.sub(r'\D', '', s)
    if digits:
        return float(digits)
    return float(fallback)

def split_address(full_address):
    addr = clean_text(full_address)
    if not addr:
        return "-", "", ""
    parts = [p.strip() for p in addr.split(",") if p.strip()]
    if len(parts) >= 3:
        b1 = ", ".join(parts[:-2])
        b2 = parts[-2]
        b3 = parts[-1]
    elif len(parts) == 2:
        b1 = parts[0]
        b2 = parts[1]
        b3 = ""
    else:
        b1 = addr
        b2 = ""
        b3 = ""
    return b1, b2, b3

def replace_in_paragraph(paragraph, old_text, new_text):
    if old_text in paragraph.text:
        full_text = paragraph.text.replace(old_text, new_text)
        if paragraph.runs:
            first_run = paragraph.runs[0]
            f_name = first_run.font.name
            f_size = first_run.font.size
            is_bold = first_run.font.bold
            
            first_run.text = full_text
            for r in paragraph.runs[1:]:
                r.text = ""
                
            if f_name: first_run.font.name = f_name
            if f_size: first_run.font.size = f_size
            if is_bold is not None: first_run.font.bold = is_bold
        else:
            paragraph.text = full_text

def check_maps_address(alamat, url_maps):
    """Mengecek indikasi perbedaan letak antara Teks Alamat dan Link Google Maps"""
    if not url_maps or str(url_maps).lower() == 'nan' or not alamat:
        return "➖ Data Tidak Lengkap"
    try:
        # Jika link pendek, resolve ke URL asli (Maksimal 3 detik agar tidak macet)
        if "maps.app.goo.gl" in url_maps or "goo.gl" in url_maps or "bit.ly" in url_maps:
            r = requests.get(url_maps, allow_redirects=True, timeout=3)
            full_url = r.url
        else:
            full_url = url_maps
            
        decoded_url = urllib.parse.unquote(full_url).lower().replace("+", " ")
        alamat_lower = str(alamat).lower()
        
        # Ekstrak kata kunci signifikan dari teks alamat
        words = set(re.findall(r'\b[a-z]{4,}\b', alamat_lower))
        stop_words = {"jalan", "raya", "desa", "kota", "kabupaten", "kecamatan", "provinsi", "selatan", "barat", "timur", "utara", "indonesia"}
        keywords = words - stop_words
        
        if not keywords:
            return "⚠️ Alamat terlalu singkat"
            
        # Hitung kecocokan kata kunci alamat dengan wilayah di dalam Link Maps
        match_count = sum(1 for w in keywords if w in decoded_url)
        
        if match_count == 0:
            return "❌ Indikasi Beda Letak"
        elif match_count == 1:
            return "⚠️ Cek Manual"
        else:
            return "✅ Sesuai"
    except Exception:
        return "❓ Gagal Akses Tautan"

def generate_loi_doc(template_path_or_bytes, row_data, config):
    doc = docx.Document(template_path_or_bytes)
    
    nomor_loi = clean_text(row_data.get("loi_mss", ""))
    if not nomor_loi or nomor_loi == "-":
        nomor_loi = config["nomor_loi"]
        
    tgl_surat = config["tanggal_surat"]
    tgl_persetujuan = config["tanggal_persetujuan"]
    
    landlord_raw = clean_text(row_data.get("pemilik", ""))
    landlord_upper = landlord_raw.upper() if landlord_raw else "LANDLORD"
    lokasi_nama = clean_text(row_data.get("lokasi", ""))
    
    alamat_raw = clean_text(row_data.get("alamat", ""))
    b1, b2, b3 = split_address(alamat_raw)
    
    qty_evcs = row_data.get("evcs", 0)
    qty_bss = row_data.get("bss", 0)
    
    harga_tahun = row_data.get("harga", 0)
    harga_bulan = int(harga_tahun / 12) if harga_tahun > 0 else 0
        
    for p in doc.paragraphs:
        txt_stripped = p.text.strip()
        
        if txt_stripped in [
            "LAMPIRAN I RINGKASAN KETENTUAN KOMERSIAL",
            "AWAL SURAT PENAWARAN Ref.",
            "No. LOI/JBDTBK-7436/05Okt/2026-0008",
            "LAMPIRAN II",
            "Detail Lokasi Slot Parkir",
            "Ref. No. LOI/JBDTBK-7436/05Okt/2026-0008",
            "LEMBAR PERSETUJUAN LOI"
        ]:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.left_indent = 0
            p.paragraph_format.right_indent = 0
            p.paragraph_format.first_line_indent = 0
            
        if "No. LOI/JBDTBK-7436" in p.text and p.text.startswith(" "):
            full_text = p.text.strip()
            if p.runs:
                first_run = p.runs[0]
                f_name, f_size, f_bold = first_run.font.name, first_run.font.size, first_run.font.bold
                first_run.text = full_text
                for r in p.runs[1:]: r.text = ""
                if f_name: first_run.font.name = f_name
                if f_size: first_run.font.size = f_size
                if f_bold is not None: first_run.font.bold = f_bold
            else:
                p.text = full_text

        if "Jakarta," in p.text and ("October" in p.text or "202" in p.text):
            if p.runs:
                p.runs[0].text = f"Jakarta, {tgl_surat}"
                for r in p.runs[1:]: r.text = ""
            else:
                p.text = f"Jakarta, {tgl_surat}"
        else:
            replace_in_paragraph(p, "Jakarta, 12 October 2026", f"Jakarta, {tgl_surat}")
            
        replace_in_paragraph(p, "LOI/JBDTBK-7436/05Okt/2026-0008", nomor_loi)
        replace_in_paragraph(p, "Bapak MUSLIM KARIM", f"Saudara {landlord_upper}")
        replace_in_paragraph(p, "bapak MUSLIM KARIM", f"Saudara {landlord_upper}")
        replace_in_paragraph(p, "Jalan Lintas Sumatra Dusun IV Rt.002 Rw.004 Desa Kedaton", b1)
        replace_in_paragraph(p, "Kecamatan Kalianda Kabupaten Lampung Selatan, ", (b2 + ", ") if b2 else "")
        replace_in_paragraph(p, "Propinsi Lampung ", b3)
        replace_in_paragraph(p, "Saya, MUSLIM KARIM", f"Saya, {landlord_upper}")
        replace_in_paragraph(p, "5 October 2026", tgl_persetujuan)
        
        if p.text.strip() == "MUSLIM KARIM":
            p.text = landlord_upper

    if len(doc.tables) >= 1:
        for row in doc.tables[0].rows:
            if "Jumlah Area Parkir" in row.cells[1].text:
                target_p = row.cells[2].paragraphs[0]
                f_name, f_size = None, None
                if target_p.runs:
                    f_name, f_size = target_p.runs[0].font.name, target_p.runs[0].font.size
                
                target_p.text = ""
                new_run = target_p.add_run(f"Saudara {landlord_upper} akan menyediakan lokasi aset dibawah kepemilikan Saudara {landlord_upper} dengan sistem sewa lahan per {qty_evcs} slot parking ( 2,5 x 5 m) dan {qty_bss} SPP (1 x 2 m).")
                if f_name: new_run.font.name = f_name
                if f_size: new_run.font.size = f_size
            else:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        replace_in_paragraph(p, "bapak MUSLIM KARIM", f"Saudara {landlord_upper}")
                        replace_in_paragraph(p, "MUSLIM KARIM", landlord_upper)
                        replace_in_paragraph(p, "LOI/JBDTBK-7436/05Okt/2026-0008", nomor_loi)

    if len(doc.tables) >= 2:
        t2 = doc.tables[1]
        t2.autofit = False
        t2.allow_autofit = False
        
        col_widths = [Cm(1.0), Cm(6.5), Cm(2.5), Cm(1.8), Cm(3.2), Cm(3.5)]
        
        if len(t2.rows) >= 2:
            cells = t2.rows[1].cells
            cells[0].text = "1"
            cells[1].text = f"{lokasi_nama.upper()} ({landlord_upper})"
            cells[2].text = str(qty_evcs) if qty_evcs > 0 else "-"
            cells[3].text = str(qty_bss) if qty_bss > 0 else "-"
            cells[4].text = format_rupiah(harga_bulan)
            cells[5].text = format_rupiah(harga_tahun)
            
        for row_idx, row in enumerate(t2.rows):
            for col_idx, cell in enumerate(row.cells):
                cell.width = col_widths[col_idx] 
                for p in cell.paragraphs:
                    if col_idx != 1:
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    else:
                        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                        
                    if row_idx >= 1 and p.runs:
                        p.runs[0].font.name = "Times New Roman"
                        p.runs[0].font.size = Pt(11)

    bio = io.BytesIO()
    doc.save(bio)
    bio.seek(0)
    return bio

# --- ANTARMUKA STREAMLIT ---
st.title("⚡ Automasi Pembuatan Dokumen LOI (Excel ➡️ Word)")
st.caption("PT MSS Alpha Tech — Kemitraan Landlord EVCS & BSS")
st.markdown("---")

col_a, col_b = st.columns([1, 1], gap="large")

with col_a:
    st.subheader("📁 1. Unggah File Rekap Excel")
    excel_file = st.file_uploader("Upload file Excel daftar lahan (.xlsx / .xls)", type=["xlsx", "xls"])
    
    if os.path.exists(TEMPLATE_FILE):
        template_source = TEMPLATE_FILE
        st.success(f" Template aktif: `{TEMPLATE_FILE}`")
    else:
        st.warning("Template default tidak ditemukan di server.")
        template_source = st.file_uploader("Upload Template Dokumen Word (.docx):", type=["docx"])

with col_b:
    st.subheader("⚙️ 2. Pengaturan Tanggal & Nomor Cadangan")
    now = datetime.today()
    default_date_str = now.strftime("%d %B %Y")
    
    c1, c2 = st.columns(2)
    with c1:
        tgl_surat = st.text_input("1. Tanggal Surat Terbit", value=default_date_str)
    with c2:
        tgl_persetujuan = st.text_input("2. Tanggal Lembar Persetujuan", value=default_date_str)
        
    prefix_nomor = st.text_input("Nomor Surat Cadangan (Jika kolom 'LOI MSS' di Excel kosong)", value=f"LOI/MSS-AT/{now.strftime('%d%b')}/{now.year}")
    fallback_harga_tahun = st.number_input("Harga Sewa / Tahun (Jika kolom harga Excel kosong/tidak terbaca)", value=20000000, step=1000000)

st.markdown("---")

if excel_file and template_source:
    st.subheader("📋 3. Data Terbaca dari Excel")
    try:
        xls = pd.ExcelFile(excel_file)
        sheet_name = st.selectbox("Pilih Lembar Kerja (Sheet):", xls.sheet_names, index=0)
        df_raw = pd.read_excel(excel_file, sheet_name=sheet_name, header=None)
        
        header_row_idx = -1
        for i in range(min(25, len(df_raw))):
            joined = " ".join([str(v) for v in df_raw.iloc[i].values if pd.notna(v)]).upper()
            if "NAMA LOKASI" in joined or "PEMILIK" in joined:
                header_row_idx = i
                break
                
        if header_row_idx != -1:
            header = df_raw.iloc[header_row_idx]
            df_body = df_raw.iloc[header_row_idx + 1:].copy()
            
            col_map = {"lokasi": 2, "pemilik": 3, "kabupaten": 8, "alamat": 9, "maps": 10, "evcs": 14, "bss": 16, "harga": 18, "loi_mss": -1}
            start_idx = max(0, header_row_idx - 1)
            end_idx = min(len(df_raw), header_row_idx + 2)
            header_range = df_raw.iloc[start_idx:end_idx]
            
            for col_idx in range(df_raw.shape[1]):
                col_text = " ".join([str(x).upper() for x in header_range[col_idx] if pd.notna(x)])
                if "NAMA LOKASI" in col_text: col_map["lokasi"] = col_idx
                elif "PEMILIK" in col_text: col_map["pemilik"] = col_idx
                elif "KABUPATEN" in col_text or "KOTA" in col_text: col_map["kabupaten"] = col_idx
                elif "ALAMAT" in col_text: col_map["alamat"] = col_idx
                elif "MAPS" in col_text: col_map["maps"] = col_idx
                elif "EVCS" in col_text: col_map["evcs"] = col_idx
                elif "BSS" in col_text: col_map["bss"] = col_idx
                elif "HARGA" in col_text: col_map["harga"] = col_idx
                elif "LOI MSS" in col_text or ("NO LOI" in col_text and "MSS" in col_text): col_map["loi_mss"] = col_idx

            parsed_list = []
            for _, r in df_body.iterrows():
                lok = r.iloc[col_map["lokasi"]] if col_map["lokasi"] < len(r) else None
                pem = r.iloc[col_map["pemilik"]] if col_map["pemilik"] < len(r) else None
                
                if pd.notna(lok) and pd.notna(pem) and str(lok).strip() != "" and str(lok).strip().upper() != "NAMA LOKASI":
                    alm = r.iloc[col_map["alamat"]] if col_map["alamat"] < len(r) else ""
                    kab = r.iloc[col_map["kabupaten"]] if col_map["kabupaten"] < len(r) else ""
                    ev = pd.to_numeric(r.iloc[col_map["evcs"]], errors='coerce') if col_map["evcs"] < len(r) else 0
                    bs = pd.to_numeric(r.iloc[col_map["bss"]], errors='coerce') if col_map["bss"] < len(r) else 0
                    
                    hg_val = r.iloc[col_map["harga"]] if col_map["harga"] < len(r) else fallback_harga_tahun
                    hg = extract_price(hg_val, fallback_harga_tahun)
                    
                    loi_val = ""
                    if col_map["loi_mss"] != -1 and col_map["loi_mss"] < len(r):
                        loi_raw = r.iloc[col_map["loi_mss"]]
                        loi_val = str(loi_raw).strip() if pd.notna(loi_raw) else ""
                        
                    maps_val = ""
                    if "maps" in col_map and col_map["maps"] != -1 and col_map["maps"] < len(r):
                        m_raw = r.iloc[col_map["maps"]]
                        maps_val = str(m_raw).strip() if pd.notna(m_raw) else ""
                        
                    parsed_list.append({
                        "id": len(parsed_list) + 1,
                        "lokasi": str(lok).strip(),
                        "pemilik": str(pem).strip(),
                        "alamat": str(alm).strip() if pd.notna(alm) else "",
                        "kabupaten": str(kab).strip() if pd.notna(kab) else "",
                        "maps": maps_val if maps_val and maps_val.lower() != "nan" else "",
                        "evcs": int(ev) if pd.notna(ev) else 0,
                        "bss": int(bs) if pd.notna(bs) else 0,
                        "harga": float(hg),
                        "loi_mss": loi_val if loi_val and loi_val.lower() != "nan" else ""
                    })
                    
            df_display = pd.DataFrame(parsed_list)
            df_preview = df_display[["id", "loi_mss", "lokasi", "pemilik", "evcs", "bss"]].copy()
            df_preview["Harga (Rp)"] = df_display["harga"].apply(format_rupiah)
            
            st.dataframe(df_preview, use_container_width=True, height=250)
            st.info(f"📍 Terdeteksi **{len(df_display)} titik lokasi** siap dicetak ke Word.")
            
            # --- FITUR VALIDASI MAPS ---
            with st.expander("🔍 Fitur Validasi: Cek Indikasi Perbedaan Titik Maps & Alamat", expanded=True):
                st.write("Sistem akan menelusuri tautan Google Maps (Kolom K) dan mencocokkannya dengan teks Alamat (Kolom J).")
                st.warning("⏳ **Catatan:** Jangan *refresh* browser atau klik menu lain saat proses berjalan.")
                
                # Inisialisasi Session State agar tabel tidak hilang saat loading selesai
                if "hasil_validasi" not in st.session_state:
                    st.session_state.hasil_validasi = None
                
                if st.button("🚀 Mulai Validasi Titik Maps"):
                    progress_bar = st.progress(0)
                    status_text = st.empty()
                    val_results = []
                    
                    for idx, item in enumerate(parsed_list):
                        status_text.text(f"Memvalidasi baris {idx+1} dari {len(parsed_list)} (Lokasi: {item['lokasi']})...")
                        status = check_maps_address(item['alamat'], item['maps'])
                        val_results.append({
                            "No": item['id'],
                            "Nama Lokasi": item['lokasi'],
                            "Alamat (Kolom J)": item['alamat'],
                            "Indikasi": status,
                            "Tautan Maps": item['maps']
                        })
                        progress_bar.progress((idx + 1) / len(parsed_list))
                        
                    status_text.success("✅ Validasi selesai!")
                    # Simpan hasil ke session_state agar tampil permanen
                    st.session_state.hasil_validasi = pd.DataFrame(val_results)
                
                # Menampilkan Tabel jika hasil sudah ada di dalam sesi
                if st.session_state.hasil_validasi is not None:
                    st.dataframe(st.session_state.hasil_validasi, use_container_width=True)
                    
                    # Tombol unduh hasil validasi ke Excel/CSV
                    csv_data = st.session_state.hasil_validasi.to_csv(index=False).encode('utf-8')
                    st.download_button(
                        label="📥 Unduh Hasil Validasi (.CSV)",
                        data=csv_data,
                        file_name=f"Hasil_Validasi_Maps_{now.strftime('%Y%m%d')}.csv",
                        mime="text/csv"
                    )
            
            # --- 4. EXPORT WORD ---
            st.markdown("### 📥 4. Export Dokumen MS Word (.docx)")
            tab_one, tab_all = st.tabs(["📄 Cetak Satu Dokumen (Pilihan)", "📦 Cetak Semua Dokumen (.ZIP)"])
            
            with tab_one:
                col_s1, col_s2 = st.columns([2, 1])
                with col_s1:
                    opsi_lokasi = [f"{item['id']}. {item['lokasi']} - Saudara {item['pemilik']}" for item in parsed_list]
                    terpilih = st.selectbox("Pilih Landlord / Titik Lokasi:", opsi_lokasi)
                    target_id = int(terpilih.split(".")[0])
                    baris_terpilih = next(item for item in parsed_list if item["id"] == target_id)
                    no_surat_satuan = baris_terpilih['loi_mss'] if baris_terpilih['loi_mss'] else f"{prefix_nomor}-{str(target_id).zfill(4)}"
                    st.caption(f"Nomor Surat yang Digunakan: **`{no_surat_satuan}`**")
                with col_s2:
                    st.write("")
                    st.write("")
                    cfg_single = {
                        "nomor_loi": no_surat_satuan,
                        "tanggal_surat": tgl_surat,
                        "tanggal_persetujuan": tgl_persetujuan
                    }
                    word_stream = generate_loi_doc(template_source, baris_terpilih, cfg_single)
                    nama_file_word = f"LOI_{re.sub(r'[^a-zA-Z0-9]', '_', baris_terpilih['lokasi'])}_{re.sub(r'[^a-zA-Z0-9]', '_', baris_terpilih['pemilik'])}.docx"
                    
                    st.download_button(
                        label=f"📥 Unduh Word: {baris_terpilih['lokasi']}.docx",
                        data=word_stream,
                        file_name=nama_file_word,
                        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                    )
                    
            with tab_all:
                st.write(f"Ekspor seluruh **{len(parsed_list)} dokumen LOI** ke dalam satu arsip ZIP.")
                if st.button("🚀 Proses & Unduh Seluruh Dokumen (.ZIP)"):
                    zip_mem = io.BytesIO()
                    bar = st.progress(0)
                    with zipfile.ZipFile(zip_mem, "w", zipfile.ZIP_DEFLATED) as zf:
                        for idx, item in enumerate(parsed_list):
                            no_loi_item = item['loi_mss'] if item['loi_mss'] else f"{prefix_nomor}-{str(item['id']).zfill(4)}"
                            cfg_item = {
                                "nomor_loi": no_loi_item,
                                "tanggal_surat": tgl_surat,
                                "tanggal_persetujuan": tgl_persetujuan
                            }
                            w_bytes = generate_loi_doc(template_source, item, cfg_item)
                            f_item_name = f"LOI_{str(item['id']).zfill(3)}_{re.sub(r'[^a-zA-Z0-9]', '_', item['lokasi'])}_{re.sub(r'[^a-zA-Z0-9]', '_', item['pemilik'])}.docx"
                            zf.writestr(f_item_name, w_bytes.getvalue())
                            bar.progress((idx + 1) / len(parsed_list))
                            
                    zip_mem.seek(0)
                    st.success("✅ Seluruh file MS Word berhasil dikemas!")
                    st.download_button(
                        label="📦 Unduh Arsip ZIP (.ZIP)",
                        data=zip_mem,
                        file_name=f"Batch_LOI_MSS_Alpha_Tech_{now.strftime('%Y%m%d')}.zip",
                        mime="application/zip"
                    )
        else:
            st.error("Baris judul kolom 'NAMA LOKASI' atau 'PEMILIK' tidak ditemukan di sheet ini.")
    except Exception as err:
        st.error(f"Gagal memproses file Excel: {err}")